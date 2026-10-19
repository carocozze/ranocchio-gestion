"""
Servicio de exportación a Excel (.xlsx) completo para Ranocchio
"""
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import crud

def generate_full_excel() -> io.BytesIO:
    wb = openpyxl.Workbook()
    
    # Colores institucionales Ranocchio
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    sub_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    sub_font = Font(name="Segoe UI", size=10, bold=True, color="334155")
    data_font = Font(name="Segoe UI", size=10, color="0F172A")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # 1. HOJA RESUMEN Y P&L
    ws_pnl = wb.active
    ws_pnl.title = "P&L Mensual"
    ws_pnl.views.sheetView[0].showGridLines = True

    pnl_data = crud.get_monthly_pnl_data()
    headers_pnl = [
        "Mes", "Inversión (USD)", "FX", "Inversión ($)", "Ingresos Ventas ($)", 
        "Gto Variable Prod ($)", "Gto Variable Pack ($)", "Total Gto Variable ($)",
        "Gto Fijo Showroom ($)", "Gto Fijo Admin ($)", "Gto Fijo Mkt ($)", "Gto Fijo Viáticos ($)", 
        "Total Gto Fijo ($)", "Gasto Total ($)", "Utilidad Neta ($)", "% Margen"
    ]
    for col_num, h in enumerate(headers_pnl, 1):
        cell = ws_pnl.cell(row=1, column=col_num, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r_idx, r in enumerate(pnl_data, 2):
        row_vals = [
            r["month_label"], r["investment_usd"], r["fx_rate"], r["investment_pesos"],
            r["effective_sales"], r["var_expense_product"], r["var_expense_packaging"],
            r["total_var_expense"], r["fixed_expense_showroom"], r["fixed_expense_admin"],
            r["fixed_expense_marketing"], r["fixed_expense_travel"], r["total_fixed_expense"],
            r["effective_expenses"], r["effective_net"], f"{r['effective_margin_pct']}%"
        ]
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_pnl.cell(row=r_idx, column=c_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if c_idx > 1 and c_idx < 16 and isinstance(val, (int, float)):
                cell.number_format = "$#,##0"
                cell.alignment = Alignment(horizontal="right")
            elif c_idx == 1:
                cell.alignment = Alignment(horizontal="left")
                cell.font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")

    # 2. HOJA VENTAS
    ws_sales = wb.create_sheet(title="Ventas")
    ws_sales.views.sheetView[0].showGridLines = True
    sales = crud.list_sales(limit=5000)
    headers_sales = [
        "ID", "Mes", "Fecha", "Cliente", "Canal", "Producto", "Tipo Orden", "Cantidad",
        "Precio Unitario", "Facturado", "Envío", "Total Venta", "Medio de Pago",
        "Costo Unitario", "Costo Total", "Margen Comercial $", "Margen %", "Detalles"
    ]
    for col_num, h in enumerate(headers_sales, 1):
        cell = ws_sales.cell(row=1, column=col_num, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for r_idx, s in enumerate(sales, 2):
        row_vals = [
            s["id"], s["month_label"], s["date"], s["customer"], s["channel"], s["product_name"],
            s["order_type"], s["quantity"], s["unit_price"], s["billed"], s["shipping"], s["total_sale"],
            s["payment_method"], s["unit_cost"], s["total_cost"], s["commercial_margin_amount"],
            f"{round(s['commercial_margin_pct']*100, 1)}%", s["details"]
        ]
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_sales.cell(row=r_idx, column=c_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if c_idx in [9, 10, 11, 12, 14, 15, 16] and isinstance(val, (int, float)):
                cell.number_format = "$#,##0"
                cell.alignment = Alignment(horizontal="right")

    # 3. HOJA GASTOS
    ws_expenses = wb.create_sheet(title="Gastos")
    ws_expenses.views.sheetView[0].showGridLines = True
    expenses = crud.list_expenses(limit=5000)
    headers_expenses = [
        "ID", "Mes", "Fecha", "Categoría", "Rubro / Subcategoría", "Detalle", 
        "Factura #", "Monto Pagado", "Forma de Pago", "Notas"
    ]
    for col_num, h in enumerate(headers_expenses, 1):
        cell = ws_expenses.cell(row=1, column=col_num, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for r_idx, e in enumerate(expenses, 2):
        row_vals = [
            e["id"], e["month_label"], e["date"], e["category"], e["subcategory"],
            e["detail"], e["invoice_number"], e["amount"], e["payment_method"], e["notes"]
        ]
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_expenses.cell(row=r_idx, column=c_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if c_idx == 8 and isinstance(val, (int, float)):
                cell.number_format = "$#,##0"
                cell.alignment = Alignment(horizontal="right")

    # 4. HOJA STOCK
    ws_stock = wb.create_sheet(title="Stock & Inventario")
    ws_stock.views.sheetView[0].showGridLines = True
    stock = crud.list_stock()
    headers_stock = [
        "ID", "Pieza", "Modelo", "Estampado", "Temporada", "Talle",
        "Producción", "Stock Total", "Showroom", "Tienda Nube", "Consignación", "Venta",
        "Mínimo", "Estado", "Acción", "Precio Unitario", "Valor Stock ($)", "Costo Unitario", "Costo Stock ($)"
    ]
    for col_num, h in enumerate(headers_stock, 1):
        cell = ws_stock.cell(row=1, column=col_num, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for r_idx, st in enumerate(stock, 2):
        row_vals = [
            st["id"], st["piece"], st["model"], st["print_design"], st["season"], st["size"],
            st["production_qty"], st["stock_total"], st["showroom_qty"], st["tiendanube_qty"],
            st["consignment_qty"], st["sold_qty"], st["min_stock"], st["status"], st["action"],
            st["unit_price"], st["stock_value"], st["unit_cost"], st["stock_cost"]
        ]
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_stock.cell(row=r_idx, column=c_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if c_idx in [16, 17, 18, 19] and isinstance(val, (int, float)):
                cell.number_format = "$#,##0"
                cell.alignment = Alignment(horizontal="right")

    # Ajustar ancho de columnas en todas las hojas
    for ws in [ws_pnl, ws_sales, ws_expenses, ws_stock]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 38)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
